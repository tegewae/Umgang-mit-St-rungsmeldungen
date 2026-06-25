"""
Importiert bestehende Störungsdaten aus der Dashboard-Excel-Datei
in die SQLite-Datenbank.

Aufruf: python import_excel.py <pfad_zur_excel_datei>
        python import_excel.py "Störmeldungen_2026Dashboard_v8.xlsx"
"""
import sys
import os
import sqlite3
from datetime import datetime, date

def excel_date(val):
    """Konvertiert Excel-Seriennummer oder String in ISO-Datum."""
    if not val:
        return None
    if isinstance(val, (int, float)):
        # Excel-Seriennummer → datetime
        from datetime import timedelta
        origin = datetime(1899, 12, 30)
        d = origin + timedelta(days=float(val))
        return d.strftime("%Y-%m-%d")
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    return str(val)[:10]

def excel_time(val):
    """Konvertiert Excel-Zeitfraktion in HH:MM."""
    if not val:
        return None
    if isinstance(val, float):
        total_min = round(val * 24 * 60)
        total_min = total_min % (24 * 60)
        return f"{total_min // 60:02d}:{total_min % 60:02d}"
    if isinstance(val, datetime):
        return val.strftime("%H:%M")
    return str(val)[:5]

def run(excel_path, db_path="stoerungen.db"):
    try:
        import openpyxl
    except ImportError:
        print("Fehler: openpyxl nicht installiert. Bitte 'pip install openpyxl' ausführen.")
        sys.exit(1)

    if not os.path.exists(excel_path):
        print(f"Fehler: Datei nicht gefunden: {excel_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # Sheet-Name finden
    target = None
    for name in wb.sheetnames:
        if "bersicht" in name or name.lower() in ("daten", "data", "stoerungen"):
            target = name
            break
    if not target:
        target = wb.sheetnames[0]

    ws = wb[target]
    print(f"Lese Sheet '{target}'...")

    # Header-Zeile finden (erste Zeile mit 'Störungsnummer' oder 'Nummer')
    header_row = None
    col_map    = {}
    for row in ws.iter_rows(min_row=1, max_row=20):
        for cell in row:
            v = str(cell.value or "").strip()
            if any(k in v for k in ("Störungsnummer","Nummer","nummer")):
                header_row = cell.row
                break
        if header_row:
            break

    if not header_row:
        print("Fehler: Keine Kopfzeile mit 'Störungsnummer' gefunden.")
        sys.exit(1)

    # Spalten-Mapping
    for cell in ws[header_row]:
        v = str(cell.value or "").lower().strip()
        if "nummer" in v:         col_map["nummer"]        = cell.column
        elif "aufgaben" in v:     col_map["aufgabenbereich"] = cell.column
        elif "spezifisch" in v or "ort" in v: col_map["spez_ort"] = cell.column
        elif "datum" in v:        col_map["datum_beginn"]  = cell.column
        elif "beginn" in v or ("uhrzeit" in v and "end" not in v):
            col_map["uhrzeit_beginn"] = cell.column
        elif "end" in v and "einsatz" in v: col_map["uhrzeit_ende"] = cell.column
        elif "meldetext" in v or "meldung" in v: col_map["meldetext"] = cell.column
        elif "bereitschaft" in v: col_map["bereitschaftler"] = cell.column
        elif "einsatzart" in v:   col_map["laptop_einsatz"] = cell.column
        elif "kommentar" in v:    col_map["bemerkungen"]   = cell.column

    print(f"Erkannte Spalten: {col_map}")

    db  = sqlite3.connect(db_path)
    db.execute("PRAGMA foreign_keys=ON")
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Standard-Benutzer für Import
    user = db.execute("SELECT id FROM benutzer LIMIT 1").fetchone()
    if not user:
        db.execute("INSERT INTO benutzer(name,passwort,rolle) VALUES('Import','x','bereitschaft')")
        db.commit()
        user = db.execute("SELECT id FROM benutzer WHERE name='Import'").fetchone()
    user_id = user[0]

    imported = 0
    skipped  = 0

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        nummer = str(row[col_map["nummer"] - 1] or "").strip() if "nummer" in col_map else None
        if not nummer or nummer in ("None",""):
            continue

        # Duplikat prüfen
        exists = db.execute("SELECT id FROM stoerung WHERE nummer=?", (nummer,)).fetchone()
        if exists:
            skipped += 1
            continue

        def get_col(key, default=""):
            if key not in col_map:
                return default
            v = row[col_map[key] - 1]
            return v if v is not None else default

        datum         = excel_date(get_col("datum_beginn"))
        uhrzeit_start = excel_time(get_col("uhrzeit_beginn"))
        uhrzeit_ende  = excel_time(get_col("uhrzeit_ende"))
        meldetext     = str(get_col("meldetext","")).strip()
        aufgaben      = str(get_col("aufgabenbereich","")).strip() or "Unbekannt"
        spez_ort      = str(get_col("spez_ort","")).strip()
        bereitsch     = str(get_col("bereitschaftler","")).strip()
        einsatzart    = str(get_col("laptop_einsatz","")).strip().lower()
        laptop        = 1 if "laptop" in einsatzart else 0
        bemerkungen   = str(get_col("bemerkungen","")).strip()

        if not datum or not meldetext:
            skipped += 1
            continue

        # Bereitschaftler als Benutzer anlegen / finden
        if bereitsch:
            b_user = db.execute("SELECT id FROM benutzer WHERE name=?", (bereitsch,)).fetchone()
            if not b_user:
                from werkzeug.security import generate_password_hash
                db.execute(
                    "INSERT INTO benutzer(name,passwort,rolle) VALUES(?,?,?)",
                    (bereitsch, generate_password_hash(bereitsch.lower() + "123"), "bereitschaft")
                )
                db.commit()
                b_user = db.execute("SELECT id FROM benutzer WHERE name=?", (bereitsch,)).fetchone()
            uid = b_user[0]
        else:
            uid = user_id

        db.execute("""
            INSERT INTO stoerung
              (nummer, aufgabenbereich, spez_ort, datum_beginn, uhrzeit_beginn,
               uhrzeit_ende, meldetext, bemerkungen, laptop_einsatz,
               status, erstellt_von, erstellt_am)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?)
        """, (
            nummer, aufgaben, spez_ort, datum, uhrzeit_start,
            uhrzeit_ende, meldetext, bemerkungen, laptop,
            "freigegeben", uid, now
        ))

        # Freigabe-Log für importierte Daten
        sid = db.execute("SELECT last_insert_rowid()").fetchone()[0]
        meister = db.execute("SELECT id FROM benutzer WHERE rolle='meister' LIMIT 1").fetchone()
        if meister:
            db.execute(
                "INSERT INTO freigabe_log(stoerung_id,aktion,von_user,zeitstempel,kommentar) VALUES(?,?,?,?,?)",
                (sid, "freigegeben", meister[0], now, "Import aus Excel-Datei")
            )

        imported += 1

    db.commit()
    db.close()
    print(f"\n✓ Import abgeschlossen: {imported} importiert, {skipped} übersprungen (Duplikate oder leer).")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Aufruf: python import_excel.py <excel-datei> [datenbank.db]")
        sys.exit(1)
    db_path = sys.argv[2] if len(sys.argv) > 2 else "stoerungen.db"
    run(sys.argv[1], db_path)
